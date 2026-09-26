#!/usr/bin/env python3
"""Build Alper Kulturel's CV as PDF (reportlab) and DOCX (python-docx)."""

import os

# Resolved from this file's own location so the script runs from anywhere:
# tools/build_cv.py -> <repo>/assets
OUT_DIR = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets")
STEM = "Alper-Kulturel-CV"

NAME = "Alper Kulturel"
ROLE = "Financial Analyst  ·  Data Analyst  ·  Quantitative Analyst"
CONTACT = "kulturelalper@gmail.com  ·  linkedin.com/in/alper-kulturel-951a7421a"
CONTACT2 = "github.com/Alper-Kulturel  ·  alper-kulturel.github.io  ·  Istanbul, Türkiye (open to US relocation)"

SUMMARY = (
    "Financial Analyst with 2+ years of experience in financial data analysis, performance "
    "reporting, and KPI-driven decision support across finance, AI, and technology. In addition "
    "to analytical work, I develop the tools behind it, with dozens of public projects on GitHub "
    "covering automated reporting pipelines, risk models, and trading strategy backtests. "
    "Currently seeking Financial Analyst, Data Analyst, or Quantitative Analyst roles in the US."
)

COMPETENCIES = (
    "Financial Analysis  ·  Performance Reporting  ·  Budget vs. Actuals & Variance Analysis  ·  "
    "KPI Design & Forecasting  ·  Credit Risk (PD/LGD/EAD)  ·  Risk Modelling (VaR / CVaR)  ·  "
    "Portfolio Optimisation  ·  Strategy Backtesting  ·  ETL & Data Pipelines  ·  Interpretable ML"
)

EXPERIENCE = [
    # ASCII, not "İstanbul": the PDF uses Helvetica with WinAnsi encoding,
    # which has no glyph for İ and renders it as a black box. The site keeps
    # the correct Turkish spelling.
    ("Financial Analyst", "Borsa Istanbul",
     "Istanbul, Türkiye  ·  Sep 2024 – Jul 2025", [
        "Analyzed large-scale revenue and trading activity datasets to produce monthly and "
        "quarterly financial performance reports supporting product investment and budget "
        "allocation decisions across exchange operations.",
        "Designed and executed A/B testing frameworks to assess the financial impact of market "
        "product features; presented variance analyses and ROI summaries to senior management.",
        "Supported quarterly forecasting cycles by collaborating with product, engineering, and "
        "commercial teams to translate exchange business requirements into measurable financial outputs.",
        "Built financial models and scenario analyses to evaluate the P&L impact of new product "
        "launches, listing initiatives, and cost optimization programs.",
    ]),
    ("Financial Data Analyst", "Supsis Ai",
     "Istanbul, Türkiye  ·  Apr 2023 – May 2024", [
        "Prepared, validated, and reconciled financial and operational datasets to support budget "
        "forecasts and AI model performance reports for management review.",
        "Analyzed financial metrics and cost-efficiency ratios across business units, producing "
        "actionable reports that informed strategic pricing and resource allocation decisions.",
        "Partnered with engineering and product teams to define financial acceptance criteria and "
        "measurable success metrics for new product features.",
    ]),
]

# Eighteen repository names, each with a one-line description, run to roughly
# 250pt of vertical space — which was the difference between this CV fitting on
# one page and spilling Technical Skills onto a second. The names are kept
# because they carry most of the signal (options-pricing-engine,
# credit-risk-scoring, cpp-order-book); the per-project detail lives on the
# portfolio site these repositories are written up on, which the header links.
PROJECTS = [
    ("Quantitative Finance",
     "Options pricing, event-driven backtesting, VaR/CVaR risk and portfolio optimisation.",
     ["options-pricing-engine", "stock-backtesting-engine", "pairs-trading",
      "portfolio-optimization", "monte-carlo-risk-simulator", "sentiment-analyzer",
      "market-daily"]),
    ("Financial Analytics",
     "Credit risk scoring, FP&A variance attribution and automated reporting pipelines.",
     ["credit-risk-scoring", "fpa-variance-dashboard", "financial-reporting-pipeline",
      "fpa-dashboard", "market-data-etl"]),
    ("Data Analytics",
     "Churn modelling, RFM segmentation and streaming ingestion.",
     ["banking-churn-predictor", "sales-dashboard", "crypto-data-pipeline"]),
    ("Systems & Infrastructure",
     "C++17 pricing libraries, matching engines and low-latency feed handlers.",
     ["cpp-black-scholes", "cpp-order-book", "cpp-feed-handler"]),
]

EDUCATION = [
    ("MSc, International Management & Information Systems",
     "Fachhochschule Südwestfalen, Germany"),
    ("BA, Business", "Skidmore College, USA"),
]

HONOURS = "NCAA Division III All-American, Men's Tennis Doubles (2022)"

SKILLS = [
    ("Languages", "Python, C++17, SQL, C#, JavaScript"),
    ("Data & Analytics", "pandas, NumPy, SciPy, scikit-learn, XGBoost, SHAP, statsmodels, Plotly, Streamlit, Tableau"),
    ("Data Engineering", "PostgreSQL, SQLAlchemy, Docker, WebSocket ingestion, ETL design, data validation"),
    ("Systems", "CMake, pybind11, low-latency profiling, binary protocol parsing"),
]

ACCENT = "#1E40AF"
INK = "#111111"
MUTED = "#555555"


# --------------------------------------------------------------------------- PDF

def esc(text):
    """Escape the three characters reportlab's paragraph parser treats as markup."""
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_pdf(path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_JUSTIFY
    from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph,
                                    Spacer, Table, TableStyle, KeepTogether)

    accent = colors.HexColor(ACCENT)
    ink = colors.HexColor(INK)
    muted = colors.HexColor(MUTED)
    rule = colors.HexColor("#DDDDDD")

    body = ParagraphStyle("body", fontName="Helvetica", fontSize=8.7, leading=11.1,
                          textColor=ink, alignment=TA_JUSTIFY)
    bullet = ParagraphStyle("bullet", parent=body, leftIndent=9, bulletIndent=1,
                            spaceBefore=1.2, alignment=0)
    proj = ParagraphStyle("proj", parent=body, leftIndent=9, bulletIndent=1,
                          spaceBefore=1.15, alignment=0)
    # the run of repository names under each project group — indented to sit
    # with the bullets, but without a bullet marker of its own
    names = ParagraphStyle("names", parent=body, leftIndent=9, alignment=0,
                           spaceBefore=0.9)
    sec = ParagraphStyle("sec", fontName="Helvetica-Bold", fontSize=9.6, leading=11,
                         textColor=accent, spaceBefore=4.0, spaceAfter=0.6)
    sub = ParagraphStyle("sub", fontName="Helvetica-Bold", fontSize=8.3, leading=10,
                         textColor=colors.HexColor("#333333"), spaceBefore=3.0, spaceAfter=0.4)
    job = ParagraphStyle("job", fontName="Helvetica-Bold", fontSize=8.9, leading=10.6,
                         textColor=ink, spaceBefore=4.2, spaceAfter=0.5)
    name_st = ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=21, leading=22,
                             textColor=ink, spaceAfter=1.5)
    role_st = ParagraphStyle("role", fontName="Helvetica", fontSize=9.8, leading=12.4,
                             textColor=accent, spaceAfter=2.4)
    contact_st = ParagraphStyle("contact", fontName="Helvetica", fontSize=8.0, leading=11,
                                textColor=muted)

    def section(title):
        tbl = Table([[Paragraph(title.upper(), sec)]], colWidths=[176 * mm])
        tbl.setStyle(TableStyle([
            ("LINEBELOW", (0, 0), (-1, -1), 0.7, rule),
            ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 2.4),
        ]))
        return [Spacer(1, 1.6), tbl]

    def bullets(items, style=bullet):
        # esc() matters here: reportlab parses paragraph text as mini-HTML, and
        # a raw "&" (as in "P&L") is read as the start of an entity — it came
        # out on the page as "P&L;".
        return [Paragraph(esc(t), style, bulletText="•") for t in items]

    story = []
    story.append(Paragraph(esc(NAME), name_st))
    story.append(Paragraph(esc(ROLE), role_st))
    story.append(Paragraph(esc(CONTACT), contact_st))
    story.append(Paragraph(esc(CONTACT2), contact_st))

    story += section("Professional Summary")
    story.append(Spacer(1, 3.4))
    story.append(Paragraph(esc(SUMMARY), body))

    story += section("Core Competencies")
    story.append(Spacer(1, 3.4))
    story.append(Paragraph(esc(COMPETENCIES), body))

    story += section("Experience")
    for title, org, meta, points in EXPERIENCE:
        story.append(Paragraph(
            '%s &nbsp;&middot;&nbsp; %s &nbsp;&nbsp;'
            '<font face="Helvetica" size="7.9" color="%s">%s</font>'
            % (esc(title), esc(org), MUTED, esc(meta)), job))
        story += bullets(points)

    story += section("Selected Projects")
    for group, blurb, repos in PROJECTS:
        story.append(Paragraph(
            '%s &mdash; <font color="%s">%s</font>'
            % (esc(group), MUTED, esc(blurb)), sub))
        story.append(Paragraph(
            " &nbsp;&middot;&nbsp; ".join(esc(r) for r in repos), names))

    story += section("Education")
    story.append(Spacer(1, 3.4))
    for degree, school in EDUCATION:
        story.append(Paragraph(
            '<font face="Helvetica-Bold">%s</font><br/>%s' % (esc(degree), esc(school)), proj))
    story.append(Paragraph(esc(HONOURS), proj, bulletText="•"))

    skills_block = section("Technical Skills") + [Spacer(1, 3.0)]
    for label, value in SKILLS:
        skills_block.append(Paragraph(
            '<font face="Helvetica-Bold">%s:</font> %s' % (esc(label), esc(value)), proj))
    story.append(KeepTogether(skills_block))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(muted)
        canvas.drawString(17 * mm, 9.5 * mm, "Alper Kulturel  ·  Curriculum Vitae")
        canvas.drawRightString(193 * mm, 9.5 * mm, "Page %d" % doc.page)
        canvas.setStrokeColor(rule)
        canvas.setLineWidth(0.5)
        canvas.line(17 * mm, 13 * mm, 193 * mm, 13 * mm)
        canvas.restoreState()

    doc = BaseDocTemplate(path, pagesize=A4,
                          leftMargin=17 * mm, rightMargin=17 * mm,
                          topMargin=13 * mm, bottomMargin=15 * mm,
                          title="Alper Kulturel - CV", author=NAME,
                          subject="Curriculum Vitae")
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates([PageTemplate(id="all", frames=[frame], onPage=footer)])
    doc.build(story)


# -------------------------------------------------------------------------- DOCX

def build_docx(path):
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    accent = RGBColor(0x1E, 0x40, 0xAF)
    muted = RGBColor(0x55, 0x55, 0x55)

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.7)
        s.right_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.02

    def bottom_border(par):
        pPr = par._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "2")
        bottom.set(qn("w:color"), "BBBBBB")
        borders.append(bottom)
        pPr.append(borders)

    def heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = accent
        bottom_border(p)
        return p

    def sub(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        return p

    def bullet(text, bold_lead=None):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.line_spacing = 1.02
        if bold_lead:
            r = p.add_run(bold_lead)
            r.bold = True
            p.add_run((" " if bold_lead.endswith(":") else " — ") + text)
        else:
            p.add_run(text)
        return p

    # header
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(NAME)
    r.bold = True
    r.font.size = Pt(22)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(ROLE)
    r.font.size = Pt(10)
    r.font.color.rgb = accent

    for line in (CONTACT, CONTACT2):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line)
        r.font.size = Pt(8.5)
        r.font.color.rgb = muted

    heading("Professional Summary")
    doc.add_paragraph(SUMMARY).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading("Core Competencies")
    doc.add_paragraph(COMPETENCIES).alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    heading("Experience")
    for title, org, meta, points in EXPERIENCE:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(10)
        r2 = p.add_run("  ·  " + org)
        r2.font.size = Pt(9.5)
        r3 = p.add_run("   " + meta)
        r3.font.size = Pt(8.5)
        r3.font.color.rgb = muted
        for point in points:
            bullet(point)

    heading("Selected Projects")
    for group, blurb, repos in PROJECTS:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(group + " — ")
        r.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        r2 = p.add_run(blurb)
        r2.font.size = Pt(9)
        r2.font.color.rgb = muted

        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("  ·  ".join(repos))
        r.font.size = Pt(9.5)

    heading("Education")
    for degree, school in EDUCATION:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(degree)
        r.bold = True
        p.add_run("\n" + school)
    bullet(HONOURS)

    heading("Technical Skills")
    for label, value in SKILLS:
        bullet(value, bold_lead=label + ":")

    doc.core_properties.title = "Alper Kulturel - CV"
    doc.core_properties.author = NAME
    doc.save(path)


if __name__ == "__main__":
    os.makedirs(OUT_DIR, exist_ok=True)
    pdf = os.path.join(OUT_DIR, STEM + ".pdf")
    docx = os.path.join(OUT_DIR, STEM + ".docx")
    build_pdf(pdf)
    build_docx(docx)
    for f in (pdf, docx):
        print("%-58s %7d bytes" % (f, os.path.getsize(f)))
